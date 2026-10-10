from docx import Document

from specdelta.ingest.parse_docx import (
    count_header_rows,
    doc_info,
    find_references,
    parse_spec,
    split_heading,
    table_blocks,
)


def write_sample_spec(folder):
    """A tiny spec with the structures we care about: numbered headings, a table with a
    footnote marker, an unnumbered heading, an annex and a change history annex."""
    doc = Document()
    doc.add_heading("Foreword", level=1)
    doc.add_paragraph("Front matter that should be ignored.")
    doc.add_heading("1\tScope", level=1)
    doc.add_paragraph("This document covers things, see clause 5.2 and TS 38.304 [3].")
    doc.add_heading("5\tOperating bands", level=1)
    doc.add_heading("5.2\tBand table", level=2)
    doc.add_paragraph("Table 5.2-1: NR operating bands")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Band"
    table.cell(0, 1).text = "Duplex"
    band = table.cell(1, 0).paragraphs[0]
    band.add_run("n24")
    band.add_run("16").font.superscript = True
    table.cell(1, 1).text = "FDD"
    doc.add_heading("Table 5.2-2", level=5)
    doc.styles.add_style("H6", 1)  # 3GPP style for level 6 headings
    doc.add_paragraph("5.2.0.0.1\tDeep clause", style="H6")
    doc.add_paragraph("See clause 6.5 of TS 38.331 for details.")
    doc.add_heading("Annex A (informative):\nExamples", level=8)
    doc.add_heading("A.1 General", level=1)
    doc.add_paragraph("Annex text.")
    doc.add_heading("Annex B (informative):\nChange history", level=8)
    doc.add_paragraph("This must not show up anywhere.")

    folder.mkdir()
    doc.save(folder / f"{folder.name}.docx")


def test_parse_spec(tmp_path):
    folder = tmp_path / "38300-ib0"
    write_sample_spec(folder)
    doc = parse_spec(folder)

    assert (doc.spec, doc.release, doc.version) == ("38.300", "Rel-18", "18.11.0")
    clauses = {c.clause_id: c for c in doc.clauses}
    assert list(clauses) == ["1", "5", "5.2", "5.2.0.0.1", "A", "A.1"]

    assert clauses["1"].referenced_clauses == ["5.2"]
    assert clauses["1"].referenced_specs == ["38.304"]

    band_clause = clauses["5.2"]
    assert band_clause.path == ["5 Operating bands"]
    assert "n24 (note 16) | FDD" in band_clause.text
    # The unnumbered heading stays with its clause
    assert "Table 5.2-2" in band_clause.text
    # The cross-spec clause reference is not internal
    assert clauses["5.2.0.0.1"].referenced_clauses == []

    assert clauses["A"].title == "Examples"
    assert clauses["A.1"].path == ["A Examples"]
    assert clauses["A.1"].text == "Annex text."

    assert [b["type"] for b in band_clause.blocks] == ["paragraph", "table", "paragraph"]
    tables = table_blocks(doc.clauses)
    assert len(tables) == 1
    assert tables[0]["table_id"] == "5.2/1"
    assert tables[0]["caption"] == "Table 5.2-1: NR operating bands"
    assert tables[0]["header_rows"] == 1
    assert tables[0]["rows"][1] == ["n24 (note 16)", "FDD"]

    all_text = " ".join(c.text for c in doc.clauses)
    assert "Front matter" not in all_text
    assert "must not show up" not in all_text


def test_doc_info():
    assert doc_info("38300-hh0") == ("38.300", "Rel-17", "17.17.0")
    assert doc_info("38101-1-if0") == ("38.101-1", "Rel-18", "18.15.0")
    assert doc_info("21905-j20") == ("21.905", "Rel-19", "19.2.0")


def test_split_heading():
    assert split_heading(3, "16.4.2\tPositioning") == ("16.4.2", "Positioning")
    assert split_heading(1, "D.1 General") == ("D.1", "General")
    assert split_heading(4, "6.2A.3.1\tA-MPR") == ("6.2A.3.1", "A-MPR")
    assert split_heading(8, "Annex B (informative): Void") == ("B", "Void")
    assert split_heading(1, "Foreword") is None
    assert split_heading(2, "A") is None


def test_find_references():
    clauses, specs = find_references(
        "see clause 5.2A.1, clauses 7 and 8, clause 6.5 of TS 38.331 [12] and TR 21.905."
    )
    assert clauses == ["5.2A.1", "7"]
    assert specs == ["21.905", "38.331"]


def test_count_header_rows():
    single = [["Band", "Duplex"], ["n1", "FDD"], ["n2", "FDD"]]
    assert count_header_rows(single) == 1

    merged_down = [
        ["Pre-coding", "Bandwidth", "Bandwidth"],
        ["Pre-coding", "20 MHz", "40 MHz"],
        ["Pre-coding", "Full", "Partial"],
        ["DFT-s-OFDM", "1", "2"],
    ]
    assert count_header_rows(merged_down) == 3

    empty_below_group = [["Modulation", "MPR", "MPR"], ["", "Edge", "Inner"], ["QPSK", "1", "2"]]
    assert count_header_rows(empty_below_group) == 2

    title_row = [["Table title", "Table title"], ["Config", "Bandwidth"], ["CA_n1", "5"]]
    assert count_header_rows(title_row) == 2

    # A group title alone is not enough, the next row here is data
    data_after_group = [
        ["Range", "Range", "Freq"],
        ["FR1", "FR1", "410 MHz"],
        ["FR2", "FR2-1", "x"],
    ]
    assert count_header_rows(data_after_group) == 1
