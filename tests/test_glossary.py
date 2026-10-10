from docx import Document
from docx.enum.style import WD_STYLE_TYPE

from specdelta.ingest.glossary import parse_glossary


def test_parse_glossary(tmp_path):
    doc = Document()
    doc.styles.add_style("EW", WD_STYLE_TYPE.PARAGRAPH)
    doc.styles.add_style("NO", WD_STYLE_TYPE.PARAGRAPH)
    doc.add_heading("3\tDefinitions, symbols and abbreviations", level=1)
    doc.add_heading("3.1\tDefinitions", level=2)
    doc.add_paragraph("For the purposes of the present document, the following terms apply:")
    doc.add_paragraph(
        "CAG Cell: a PLMN cell broadcasting at least one Closed Access Group identity."
    )
    doc.add_paragraph("NOTE: Also used for SNPN.", style="NO")
    doc.add_paragraph("Boundary IAB-node: as defined in TS 38.401 [4].")
    doc.add_heading("3.2\tAbbreviations", level=2)
    doc.add_heading("A", level=2)  # letter group heading, no clause number
    doc.add_paragraph("1x RTT\tCDMA2000 1x Radio Transmission Technology", style="EW")
    doc.add_paragraph("NCR\tNetwork-Controlled Repeater", style="EW")
    doc.add_heading("4\tGeneral", level=1)
    doc.add_paragraph("Normal text: this is not a definition.")

    folder = tmp_path / "38321-ia0"
    folder.mkdir()
    doc.save(folder / "38321-ia0.docx")
    entries = parse_glossary(folder)

    found = [(e.clause_id, e.kind, e.term, e.text) for e in entries]
    assert found == [
        (
            "3.1",
            "definition",
            "CAG Cell",
            "a PLMN cell broadcasting at least one Closed Access Group identity.\n"
            "NOTE: Also used for SNPN.",
        ),
        ("3.1", "definition", "Boundary IAB-node", "as defined in TS 38.401 [4]."),
        ("3.2", "abbreviation", "1x RTT", "CDMA2000 1x Radio Transmission Technology"),
        ("3.2", "abbreviation", "NCR", "Network-Controlled Repeater"),
    ]
    assert all((e.spec, e.release) == ("38.321", "Rel-18") for e in entries)
